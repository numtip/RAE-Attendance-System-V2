#!/usr/bin/env python3
"""FY2569 attendance integration dataset builder (read-only on every source).

Reads the 12 attested HIP monthly exports and emits, into ONE isolated gitignored directory:
  * a normalized per-employee-day attendance dataset
  * an explicit exception register (nothing is silently dropped)
  * an employee-master integration view with target-schema gaps
  * a manifest with source hashes and counts

Never writes to a source. Names are written only into the isolated, gitignored output
directory (the import target requires them); stdout carries counts only.

Usage (repo root):
  python scripts/data-recovery/build_fy2569_integration.py
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    sys.stderr.write("openpyxl is required\n")
    raise

MONTH_KEYS = ['oct', 'nov', 'dec', 'jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep']
KEY_TO_MONTH = {
    'oct': '2025-10', 'nov': '2025-11', 'dec': '2025-12', 'jan': '2026-01',
    'feb': '2026-02', 'mar': '2026-03', 'apr': '2026-04', 'may': '2026-05',
    'jun': '2026-06', 'jul': '2026-07', 'aug': '2026-08', 'sep': '2026-09',
}
BUDGET_YEAR_BE = 2569

C_CODE, C_WHEN = 0, 3
COL_STATUS = 'สถานะ'
COL_IN = 'เข้า [1]'
COL_OUT = 'ออก [1]'
COL_LATE = 'สาย [1]'
COL_EARLY = 'ออกก่อน [1]'

# ---- domain vocabulary (explicit and documented, never guessed per row) ------
LEAVE_MAP = {
    'ป่วย': 'sick',
    'ป่วย (ผ่าตัดนิ้วเท้า)': 'sick',
    'กิจ': 'personal',
    'พักผ่อน': 'vacation',
    'ไปราชการ': 'other',
}
OFFICIAL_DUTY = {'ไปราชการ', 'ช่วยงานกีฬา', 'งานรับปริญญา'}
ABSENT_REMARK = 'ขาดงาน'
INCOMPLETE_SCAN = {'ไม่สแกนเข้า', 'ไม่สแกนออก'}

WORKBOOK = Path('docs/RAE_Employee_Roster_Update20261020.xlsx')
IDCARD_CSV = Path('database/IDCardRaecsv2027.csv')
PREPARED = Path(r'G:\ProjectAI\attendance-report-generator\CodexWorker\prepared')
SIBLING_ROOT = Path(r'G:\ProjectAI\attendance-report-generator')
FEB_XLSX = Path(r'D:\Doc\fingerscan\2026\Total\febcsv2026xlsx.xlsx')
APR_CSV = Path(r'D:\Doc\fingerscan\2026\Total\ai\4april.csv')

DATASET_COLUMNS = [
    'employee_code', 'device_code', 'attendance_date', 'budget_year_be', 'month_be',
    'check_in', 'check_out', 'late_raw', 'early_raw', 'late_minutes', 'early_minutes',
    'is_late', 'is_early_leave', 'is_working_day', 'source_status', 'source_remark',
    'attendance_status', 'leave_type', 'classification', 'confidence', 'flags',
    'workday_source', 'source_key', 'source_file', 'source_sha256', 'source_row',
]
EXCEPTION_COLUMNS = [
    'employee_code', 'attendance_date', 'exception_type', 'detail', 'confidence',
    'source_key', 'source_file', 'source_row',
]
MASTER_COLUMNS = [
    'report_code', 'employee_code', 'device_code', 'title_th', 'full_name_th',
    'department_current', 'employment_status', 'identity_kind', 'in_master_workbook',
    'in_attendance_evidence', 'has_2027_idcard', 'first_attendance_date',
    'last_attendance_date', 'attendance_rows', 'schema_gaps',
]


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def decode_bytes(raw: bytes) -> tuple[str, str]:
    if raw.startswith(b'\xef\xbb\xbf'):
        return raw.decode('utf-8-sig'), 'utf-8-sig(BOM)'
    try:
        return raw.decode('utf-8'), 'utf-8'
    except UnicodeDecodeError:
        return raw.decode('cp874', errors='replace'), 'cp874(fallback)'


def normalize_date(value: str) -> date | None:
    """Accept DD/MM/BBBB (BE) and ISO 'BBBB-MM-DD'. Returns a Gregorian date."""
    text = str(value).strip()
    if not text:
        return None
    iso = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', text)
    if iso:
        y, m, d = (int(x) for x in iso.groups())
    else:
        slash = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$', text)
        if not slash:
            return None
        d, m, y = (int(x) for x in slash.groups())
    if y > 2400:
        y -= 543
    try:
        return date(y, m, d)
    except ValueError:
        return None


def parse_minutes(value: str) -> int:
    text = str(value).strip()
    if not text or text in ('0', '00:00:00', 'None'):
        return 0
    m = re.match(r'^(\d{1,2}):(\d{2})(?::(\d{2}))?$', text)
    return int(m.group(1)) * 60 + int(m.group(2)) if m else 0


def fmt_time(value: str) -> str:
    text = str(value).strip()
    return '' if text in ('', '0', 'None') else text


def column_index(header: list[str], name: str) -> int | None:
    return header.index(name) if name in header else None


def remark_index(header: list[str]) -> int:
    for i, h in enumerate(header):
        if h.startswith('Remark'):
            return i
    raise SystemExit('FATAL: Remark column missing')


def load_month(key: str, sources: dict) -> tuple[list[str], list[list[str]], dict]:
    entry = sources[key]
    if key == 'feb':
        path = FEB_XLSX
    elif key == 'apr':
        path = APR_CSV
    else:
        path = SIBLING_ROOT / 'input' / 'logs' / Path(entry['file']).name
    raw = path.read_bytes()
    digest = sha256_bytes(raw)
    if digest != entry['sha256']:
        raise SystemExit(f'FATAL: {path} hash mismatch (attested {entry["sha256"]})')
    if path.suffix.lower() == '.xlsx':
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        try:
            table = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
        finally:
            wb.close()
        header = [str(c).strip() if c is not None else '' for c in table[0]]
        rows = [['' if c is None else str(c).strip() for c in r] for r in table[1:]]
        rows = [r for r in rows if r and r[0].strip()]
        encoding = 'xlsx'
    else:
        text, encoding = decode_bytes(raw)
        table = list(csv.reader(io.StringIO(text)))
        header = [h.strip() for h in table[0]]
        rows = [r for r in table[1:] if r and r[0].strip()]
    prov = {'key': key, 'month': KEY_TO_MONTH[key], 'file': str(path), 'sha256': digest,
            'rows': len(rows), 'attested_rows': entry['rows_used'], 'encoding': encoding}
    return header, rows, prov


def learn_holiday_labels(loaded) -> set[str]:
    """Learn holiday labels ONLY from exports that actually carry a สถานะ column."""
    labels: set[str] = set()
    for _key, header, rows, _prov in loaded:
        si = column_index(header, COL_STATUS)
        if si is None:
            continue
        ri = remark_index(header)
        for r in rows:
            if si < len(r) and r[si].strip() == 'วันหยุด' and ri < len(r) and r[ri].strip():
                labels.add(r[ri].strip())
    return labels


def classify(row: list[str], header: list[str], holiday_labels: set[str],
             has_status: bool, weekday: int) -> dict:
    si = column_index(header, COL_STATUS)
    ri = remark_index(header)
    ii, oi = column_index(header, COL_IN), column_index(header, COL_OUT)
    li, ei = column_index(header, COL_LATE), column_index(header, COL_EARLY)

    status_raw = row[si].strip() if si is not None and si < len(row) else ''
    remark = row[ri].strip() if ri < len(row) else ''
    raw_in = row[ii].strip() if ii is not None and ii < len(row) else ''
    raw_out = row[oi].strip() if oi is not None and oi < len(row) else ''
    late_min = parse_minutes(row[li]) if li is not None and li < len(row) else 0
    early_min = parse_minutes(row[ei]) if ei is not None and ei < len(row) else 0

    flags: list[str] = []
    leave_type = ''

    if has_status:
        is_working = status_raw == 'วันทำงาน'
        if status_raw not in ('วันทำงาน', 'วันหยุด'):
            flags.append('STATUS_UNRECOGNIZED')
            is_working = bool(raw_in or raw_out or late_min or early_min)
    else:
        # These exports carry no สถานะ column. Non-working days are recovered from two facts
        # proven on the 9 exports that DO carry it: (a) no workday ever falls on Sat/Sun,
        # (b) public holidays carry a holiday label in Remark. Nothing else is assumed.
        is_weekend = weekday >= 5
        is_labelled_holiday = remark in holiday_labels and remark != ''
        is_working = not (is_weekend or is_labelled_holiday)
        flags.append('WORKDAY_INFERRED_FROM_CALENDAR_AND_HOLIDAY_VOCAB')
        if is_weekend:
            flags.append('WEEKEND_INFERRED')

    if not is_working:
        if remark:
            flags.append('HOLIDAY_LABEL')
        return {'attendance_status': 'holiday', 'leave_type': '', 'is_late': 0,
                'is_early_leave': 0, 'is_working_day': 0, 'classification': 'holiday',
                'confidence': 'DERIVED' if has_status else 'INFERRED',
                'flags': ';'.join(flags), 'late_minutes': 0, 'early_minutes': 0,
                'check_in': '', 'check_out': '', 'late_raw': '', 'early_raw': ''}

    if remark == ABSENT_REMARK:
        return {'attendance_status': 'absent', 'leave_type': '', 'is_late': 0,
                'is_early_leave': 0, 'is_working_day': 1, 'classification': 'absent',
                'confidence': 'DERIVED', 'flags': ';'.join(flags + ['ABSENT_STATED_BY_SOURCE']),
                'late_minutes': 0, 'early_minutes': 0, 'check_in': fmt_time(raw_in),
                'check_out': fmt_time(raw_out), 'late_raw': '', 'early_raw': ''}

    if remark in LEAVE_MAP:
        if not (raw_in or raw_out):
            flags.append('LEAVE_NO_SCAN')
        return {'attendance_status': 'leave', 'leave_type': LEAVE_MAP[remark], 'is_late': 0,
                'is_early_leave': 0, 'is_working_day': 1, 'classification': 'leave',
                'confidence': 'DERIVED', 'flags': ';'.join(flags),
                'late_minutes': 0, 'early_minutes': 0, 'check_in': fmt_time(raw_in),
                'check_out': fmt_time(raw_out), 'late_raw': '', 'early_raw': ''}

    official = remark in OFFICIAL_DUTY
    if official:
        flags.append('OFFICIAL_DUTY')

    has_in, has_out = bool(raw_in), bool(raw_out)
    if late_min > 0:
        status = 'late'
    elif has_in or has_out or official:
        status = 'present'
    else:
        status = 'absent'
        flags.append('ABSENT_INFERRED_NO_SCAN')

    if status != 'absent':
        if not has_in:
            flags.append('NO_SCAN_IN')
        if not has_out:
            flags.append('NO_SCAN_OUT')
    if remark in INCOMPLETE_SCAN:
        flags.append('INCOMPLETE_SCAN_REMARK')

    confidence = 'DERIVED' if (has_in or has_out or official) else 'INFERRED'
    return {'attendance_status': status, 'leave_type': leave_type,
            'is_late': 1 if late_min > 0 else 0,
            'is_early_leave': 1 if early_min > 0 else 0, 'is_working_day': 1,
            'classification': 'official_duty' if official else status,
            'confidence': confidence, 'flags': ';'.join(flags),
            'late_minutes': late_min, 'early_minutes': early_min,
            'check_in': fmt_time(raw_in), 'check_out': fmt_time(raw_out),
            'late_raw': fmt_time(row[li]) if li is not None and li < len(row) else '',
            'early_raw': fmt_time(row[ei]) if ei is not None and ei < len(row) else ''}


def read_master() -> dict[str, dict]:
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    finally:
        wb.close()
    master = {}
    for r in rows:
        code = str(int(r[0])) if isinstance(r[0], (int, float)) else str(r[0]).strip()
        owner = r[5]
        if owner is None or str(owner).strip() == '':
            owner_code = ''
        elif isinstance(owner, (int, float)):
            owner_code = str(int(owner))
        else:
            owner_code = str(owner).strip()
        resigned = str(r[7] or '').strip() in ('ลาออก', 'เกษียณอายุ/ลาออก')
        master[code] = {
            'report_code': code, 'employee_code': owner_code,
            # device_code is evidence-driven: assigned in build() only when the export set
            # actually contains rows for this code.
            'device_code': code,
            'title_th': str(r[9] or '').strip(),
            'full_name_th': str(r[4] or '').strip(),
            'department_current': str(r[6] or '').strip(),
            'employment_status': 'resigned' if resigned else 'active',
        }
    return master


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with open(path, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=columns, extrasaction='ignore')
        w.writeheader()
        for r in rows:
            w.writerow(r)


def build(out: Path, master: dict, sources: dict) -> dict:
    out.mkdir(parents=True, exist_ok=True)

    loaded, provenance = [], []
    for key in MONTH_KEYS:
        header, rows, prov = load_month(key, sources)
        loaded.append((key, header, rows, prov))
        provenance.append(prov)

    holiday_labels = learn_holiday_labels(loaded)

    dataset: list[dict] = []
    exceptions: list[dict] = []
    seen: dict[tuple[str, str], dict] = {}
    duplicates: list[dict] = []
    per_month_codes: dict[str, set[str]] = {}

    for key, header, rows, prov in loaded:
        has_status = column_index(header, COL_STATUS) is not None
        month = KEY_TO_MONTH[key]
        codes_month: set[str] = set()
        si = column_index(header, COL_STATUS)
        ri = remark_index(header)
        for offset, row in enumerate(rows, start=2):
            code = row[C_CODE].strip() if C_CODE < len(row) else ''
            if not code:
                exceptions.append({'employee_code': '', 'attendance_date': '',
                                   'exception_type': 'ROW_WITHOUT_CODE',
                                   'detail': 'source row has no device code',
                                   'confidence': 'UNRESOLVED', 'source_key': key,
                                   'source_file': prov['file'], 'source_row': offset})
                continue
            codes_month.add(code)
            when = normalize_date(row[C_WHEN] if C_WHEN < len(row) else '')
            if when is None:
                exceptions.append({'employee_code': code,
                                   'attendance_date': (row[C_WHEN] if C_WHEN < len(row) else '')[:19],
                                   'exception_type': 'UNPARSABLE_DATE',
                                   'detail': 'date could not be normalized',
                                   'confidence': 'UNRESOLVED', 'source_key': key,
                                   'source_file': prov['file'], 'source_row': offset})
                continue
            cls = classify(row, header, holiday_labels, has_status, when.weekday())
            record = {
                'employee_code': code, 'device_code': code,
                'attendance_date': when.isoformat(), 'budget_year_be': BUDGET_YEAR_BE,
                'month_be': when.month, 'check_in': cls['check_in'],
                'check_out': cls['check_out'], 'late_raw': cls['late_raw'],
                'early_raw': cls['early_raw'], 'late_minutes': cls['late_minutes'],
                'early_minutes': cls['early_minutes'], 'is_late': cls['is_late'],
                'is_early_leave': cls['is_early_leave'],
                'is_working_day': cls['is_working_day'],
                'source_status': row[si].strip() if si is not None and si < len(row) else '',
                'source_remark': row[ri].strip() if ri < len(row) else '',
                'attendance_status': cls['attendance_status'],
                'leave_type': cls['leave_type'], 'classification': cls['classification'],
                'confidence': cls['confidence'], 'flags': cls['flags'],
                'workday_source': 'source_status' if has_status else 'inferred',
                'source_key': key, 'source_file': prov['file'],
                'source_sha256': prov['sha256'], 'source_row': offset,
            }
            dedup = (code, record['attendance_date'])
            if dedup in seen:
                duplicates.append({'employee_code': code,
                                   'attendance_date': record['attendance_date'],
                                   'first_source': seen[dedup]['source_key'],
                                   'second_source': key})
                exceptions.append({'employee_code': code,
                                   'attendance_date': record['attendance_date'],
                                   'exception_type': 'DUPLICATE_EMPLOYEE_DATE',
                                   'detail': f"already present from {seen[dedup]['source_key']}",
                                   'confidence': 'UNRESOLVED', 'source_key': key,
                                   'source_file': prov['file'], 'source_row': offset})
                continue
            seen[dedup] = record
            dataset.append(record)
            if cls['confidence'] == 'INFERRED':
                exceptions.append({'employee_code': code,
                                   'attendance_date': record['attendance_date'],
                                   'exception_type': 'ATTENDANCE_INFERRED',
                                   'detail': f"{cls['attendance_status']} without full scan evidence",
                                   'confidence': 'INFERRED', 'source_key': key,
                                   'source_file': prov['file'], 'source_row': offset})
        per_month_codes[month] = codes_month

    dataset.sort(key=lambda r: (r['attendance_date'], r['employee_code']))

    csv_codes = set()
    for row in list(csv.reader(io.StringIO(decode_bytes(IDCARD_CSV.read_bytes())[0])))[1:]:
        if row and len(row) > 1 and row[1].strip():
            csv_codes.add(row[1].strip())

    all_codes = sorted({r['employee_code'] for r in dataset} | set(master))
    rows_by_code: dict[str, list[str]] = {}
    for r in dataset:
        rows_by_code.setdefault(r['employee_code'], []).append(r['attendance_date'])

    master_rows = []
    for code in all_codes:
        m = master.get(code)
        dates = sorted(rows_by_code.get(code, []))
        if m is None:
            master_rows.append({'report_code': code, 'employee_code': '', 'device_code': code,
                                'in_master_workbook': 'no', 'in_attendance_evidence': 'yes',
                                'has_2027_idcard': 'yes' if code in csv_codes else 'no',
                                'first_attendance_date': dates[0] if dates else '',
                                'last_attendance_date': dates[-1] if dates else '',
                                'attendance_rows': str(len(dates)),
                                'identity_kind': 'UNRESOLVED',
                                'schema_gaps': 'NOT_IN_OWNER_MASTER'})
            continue
        gaps = ['first_name_th/last_name_th (one owner name field; split not approved)',
                'email (NOT NULL UNIQUE; no source)',
                'employee_type (NOT NULL; no source)',
                'employee_uid (allocation policy not approved)']
        if not m['employee_code']:
            gaps.append('employee_code (not supplied)')
        if m['employment_status'] == 'resigned':
            gaps.append('effective_to (exact resignation date unverified)')
        master_rows.append({
            'report_code': code, 'employee_code': m['employee_code'],
            'device_code': m['device_code'] if dates else '', 'title_th': m['title_th'],
            'full_name_th': m['full_name_th'],
            'department_current': m['department_current'],
            'employment_status': m['employment_status'], 'identity_kind': 'UNRESOLVED',
            'in_master_workbook': 'yes',
            'in_attendance_evidence': 'yes' if dates else 'no',
            'has_2027_idcard': 'yes' if code in csv_codes else 'no',
            'first_attendance_date': dates[0] if dates else '',
            'last_attendance_date': dates[-1] if dates else '',
            'attendance_rows': str(len(dates)), 'schema_gaps': '; '.join(gaps),
        })

    write_csv(out / 'daily_attendance_normalized.csv', DATASET_COLUMNS, dataset)
    write_csv(out / 'attendance_exceptions.csv', EXCEPTION_COLUMNS, exceptions)
    write_csv(out / 'employee_master_integration.csv', MASTER_COLUMNS, master_rows)

    source_total = sum(p['rows'] for p in provenance)
    flag_counts = Counter(f for r in dataset for f in r['flags'].split(';') if f)
    summary = {
        'generated_at_utc': datetime.now(timezone.utc).isoformat(timespec='seconds'),
        'owner_decision_e1': 'OWNER_CONFIRMED_DISTINCT_IDENTITIES',
        'e1_codes': ['1015', '1032'],
        'source_files': len(provenance),
        'row_accounting': {
            'source_rows': source_total,
            'normalized_rows': len(dataset),
            'duplicates_removed': len(duplicates),
            'no_silent_row_loss': source_total == len(dataset) + len(duplicates),
            'exception_rows': len(exceptions),
        },
        'per_month': {p['month']: {'key': p['key'], 'rows': p['rows'],
                                   'attested_rows': p['attested_rows'],
                                   'sha256': p['sha256'],
                                   'distinct_codes': len(per_month_codes[p['month']])}
                      for p in provenance},
        'distinct_codes_all_months': len({r['employee_code'] for r in dataset}),
        'status_counts': dict(Counter(r['attendance_status'] for r in dataset)),
        'confidence_counts': dict(Counter(r['confidence'] for r in dataset)),
        'flag_counts': dict(sorted(flag_counts.items())),
        'holiday_labels_learned': len(holiday_labels),
        'identity_separation': {
            '1015_rows': len(rows_by_code.get('1015', [])),
            '1032_rows': len(rows_by_code.get('1032', [])),
            '1015_months': sorted({d[:7] for d in rows_by_code.get('1015', [])}),
            '1032_months_count': len({d[:7] for d in rows_by_code.get('1032', [])}),
            'shared_dates': len(set(rows_by_code.get('1015', [])) & set(rows_by_code.get('1032', []))),
        },
    }
    (out / 'integration_manifest.json').write_text(
        json.dumps({'summary': summary, 'provenance': provenance}, ensure_ascii=False, indent=2) + '\n',
        encoding='utf-8')
    return summary


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--out', default='database/local/fy2569-integration-20261020')
    args = ap.parse_args(argv)

    raw = json.loads((PREPARED / '_build_meta.json').read_text(encoding='utf-8-sig'))
    sources = {s['key']: s for s in raw['sources']}
    summary = build(Path(args.out), read_master(), sources)

    sys.stdout.write(json.dumps({
        'normalized_rows': summary['row_accounting']['normalized_rows'],
        'source_rows': summary['row_accounting']['source_rows'],
        'no_silent_row_loss': summary['row_accounting']['no_silent_row_loss'],
        'duplicates': summary['row_accounting']['duplicates_removed'],
        'distinct_codes': summary['distinct_codes_all_months'],
        'status_counts': summary['status_counts'],
        'confidence_counts': summary['confidence_counts'],
        '1015': summary['identity_separation']['1015_rows'],
        '1032': summary['identity_separation']['1032_rows'],
    }, ensure_ascii=False) + '\n')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
